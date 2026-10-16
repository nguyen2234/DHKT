"""
Excel Report Exporter for THFT Report Tool.
Generates a multi-worksheet workbook where every center has its own dedicated sheet,
accurately mimicking the layout, colors, typography, borders, and print configuration
of the Word template.
"""

import os
import re
import datetime
import logging
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import Dict, Any, List, Optional

from .models import EmployeeFT, CenterReport, ReportDataset
from .utils import sanitize_sheet_name

logger = logging.getLogger(__name__)

# Exact Hex Colors extracted from Word Template
COLOR_GREEN = "92D050"       # A-D & Q-S
COLOR_LIGHT_OLIVE = "EBF1DE" # E-G (Bảo dưỡng nhà trạm)
COLOR_PEACH = "FDE9D9"       # H-J (Bảo dưỡng THC)
COLOR_AQUA = "92CDDC"        # K-M (Tuần tra tuyến cáp)
COLOR_DARK_BLUE = "0070C0"   # N-P (Báo hỏng hạ tầng mạng lưới)
COLOR_RED = "FF0000"         # T (Thực hiện ngày)
COLOR_NOTICE_BG = "FFF9E6"   # Soft notification card background
COLOR_NOTICE_BORDER = "F0C36D"
COLOR_TOTAL_BG = "E2EFDA"    # Soft green for summary row


class ReportExporter:
    def __init__(self, config: Dict[str, Any]):
        self.config = config
        self.font_family = config.get("export", {}).get("font_family", "Times New Roman")

    def export_excel(
        self,
        dataset: ReportDataset,
        notification_texts: Dict[str, str], # center_code -> customized notification text
        output_dir: Optional[str] = None,
        position_filter: str = "ALL",
        progress_callback=None
    ) -> str:
        """
        Exports the multi-sheet workbook: BaoCao_THFT_YYYYMMDD.xlsx.
        Supports position filtering (e.g. NV DM, NV NHÀ TRẠM).
        Returns the absolute path to the generated Excel file.
        """
        if not output_dir:
            output_dir = os.path.dirname(dataset.file_path)

        os.makedirs(output_dir, exist_ok=True)

        # Build filename: BaoCao_THFT_YYYYMMDD[_NVDM].xlsx
        date_clean = dataset.report_date_str.replace("/", "")
        if len(date_clean) == 8:
            # DDMMYYYY -> YYYYMMDD
            ymd = f"{date_clean[4:8]}{date_clean[2:4]}{date_clean[0:2]}"
        else:
            ymd = datetime.datetime.now().strftime("%Y%m%d")

        pos_suffix = ""
        if position_filter and position_filter.upper() not in ("ALL", "TẤT CẢ", "TAT CA"):
            if "DM" in position_filter.upper():
                pos_suffix = "_NVDM"
            elif "TRẠM" in position_filter.upper() or "TRAM" in position_filter.upper():
                pos_suffix = "_NVNhaTram"
            else:
                clean_p = re.sub(r'[^A-Za-z0-9]', '', position_filter)
                pos_suffix = f"_{clean_p}"

        base_filename = f"BaoCao_THFT_{ymd}{pos_suffix}.xlsx"
        output_path = os.path.join(output_dir, base_filename)

        # Avoid collision if file is currently open/locked
        counter = 1
        while os.path.exists(output_path):
            try:
                # Test open write
                with open(output_path, "a"):
                    pass
                break
            except IOError:
                output_path = os.path.join(output_dir, f"BaoCao_THFT_{ymd}_{counter}.xlsx")
                counter += 1

        wb = openpyxl.Workbook()
        # Remove default sheet
        wb.remove(wb.active)

        total_centers = len(dataset.centers)
        current_idx = 0

        # Create a worksheet for each center
        for center_code, center in dataset.centers.items():
            current_idx += 1
            if progress_callback:
                pct = int((current_idx / total_centers) * 90)
                progress_callback(pct, f"Đang tạo trang báo cáo cho {center_code} ({center.center_name})...")

            sheet_title = sanitize_sheet_name(center_code)
            ws = wb.create_sheet(title=sheet_title)

            notification_text = notification_texts.get(
                center_code,
                f"FT tồn quá hạn wo đến ngày hôm nay {dataset.day_month_str}, khẩn trương hoàn thành trước {dataset.deadline_str}. FT không có tiến độ ngày sẽ triệu tập họp vào {dataset.meeting_time_str} ngày mai\n\n{center.build_tag_string()}"
            )

            self._build_center_sheet(ws, center, notification_text)

        if progress_callback:
            progress_callback(95, "Đang lưu file Excel...")

        wb.save(output_path)

        if progress_callback:
            progress_callback(100, f"Đã xuất thành công: {os.path.basename(output_path)}")

        logger.info(f"Successfully exported report to: {output_path}")
        return output_path

    def _build_center_sheet(self, ws, center: CenterReport, notification_text: str):
        """Constructs one complete sheet matching the Word template layout."""
        # 1. Styles setup
        thin_border = Border(
            left=Side(style="thin", color="A6A6A6"),
            right=Side(style="thin", color="A6A6A6"),
            top=Side(style="thin", color="A6A6A6"),
            bottom=Side(style="thin", color="A6A6A6")
        )
        double_bottom_border = Border(
            left=Side(style="thin", color="A6A6A6"),
            right=Side(style="thin", color="A6A6A6"),
            top=Side(style="thin", color="A6A6A6"),
            bottom=Side(style="double", color="333333")
        )

        font_header_black = Font(name=self.font_family, size=10, bold=True, color="000000")
        font_header_white = Font(name=self.font_family, size=10, bold=True, color="FFFFFF")
        font_data = Font(name=self.font_family, size=10, color="000000")
        font_data_bold = Font(name=self.font_family, size=10, bold=True, color="000000")
        font_data_tagged = Font(name=self.font_family, size=10, bold=True, color="C00000") # Red highlight for tagged
        font_notice = Font(name=self.font_family, size=11, bold=True, color="1F4E78")

        align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        align_left = Alignment(horizontal="left", vertical="center")
        align_notice = Alignment(horizontal="left", vertical="center", wrap_text=True)

        fill_green = PatternFill(start_color=COLOR_GREEN, end_color=COLOR_GREEN, fill_type="solid")
        fill_olive = PatternFill(start_color=COLOR_LIGHT_OLIVE, end_color=COLOR_LIGHT_OLIVE, fill_type="solid")
        fill_peach = PatternFill(start_color=COLOR_PEACH, end_color=COLOR_PEACH, fill_type="solid")
        fill_aqua = PatternFill(start_color=COLOR_AQUA, end_color=COLOR_AQUA, fill_type="solid")
        fill_dark_blue = PatternFill(start_color=COLOR_DARK_BLUE, end_color=COLOR_DARK_BLUE, fill_type="solid")
        fill_red = PatternFill(start_color=COLOR_RED, end_color=COLOR_RED, fill_type="solid")
        fill_notice = PatternFill(start_color=COLOR_NOTICE_BG, end_color=COLOR_NOTICE_BG, fill_type="solid")
        fill_total = PatternFill(start_color=COLOR_TOTAL_BG, end_color=COLOR_TOTAL_BG, fill_type="solid")

        # 2. TOP SECTION: Notification Box (Row 1 & 2)
        # Calculate dynamic height based on text length
        lines = notification_text.count("\n") + int(len(notification_text) / 120) + 1
        ws.row_dimensions[1].height = max(45, min(120, lines * 20))

        ws.merge_cells("A1:T1")
        cell_n1 = ws["A1"]
        cell_n1.value = notification_text
        cell_n1.font = font_notice
        cell_n1.fill = fill_notice
        cell_n1.alignment = align_notice

        # Style all cells in merged range A1:T1 for clean borders
        for col_idx in range(1, 21):
            c = ws.cell(row=1, column=col_idx)
            c.border = Border(
                top=Side(style="medium", color=COLOR_NOTICE_BORDER),
                bottom=Side(style="medium", color=COLOR_NOTICE_BORDER),
                left=Side(style="medium", color=COLOR_NOTICE_BORDER) if col_idx == 1 else None,
                right=Side(style="medium", color=COLOR_NOTICE_BORDER) if col_idx == 20 else None,
            )

        ws.row_dimensions[2].height = 10 # Spacing row

        # 3. TABLE HEADERS (Row 3 & Row 4)
        h1_row = 3
        h2_row = 4
        ws.row_dimensions[h1_row].height = 24
        ws.row_dimensions[h2_row].height = 22

        # Header 1 labels
        headers_row1 = [
            ("FT", 1, 1, fill_green, font_header_black),
            ("Tên FT", 2, 2, fill_green, font_header_black),
            ("TTH", 3, 3, fill_green, font_header_black),
            ("Vị trí", 4, 4, fill_green, font_header_black),
            ("ICMS_Bảo dưỡng nhà trạm", 5, 7, fill_olive, font_header_black),
            ("ICMS_Bảo dưỡng THC", 8, 10, fill_peach, font_header_black),
            ("ICMS_Tuần tra tuyến cáp", 11, 13, fill_aqua, font_header_black),
            ("Báo hỏng hạ tầng mạng lưới", 14, 16, fill_dark_blue, font_header_white),
            ("Tổng", 17, 19, fill_green, font_header_black),
            ("Thực hiện ngày", 20, 20, fill_red, font_header_white)
        ]

        # Apply Header 1 merges and text
        for text, c_start, c_end, fill, font in headers_row1:
            if c_start == c_end and c_start in (1, 2, 3, 4, 20):
                # Vertical merge row 3 & 4
                ws.merge_cells(start_row=h1_row, start_column=c_start, end_row=h2_row, end_column=c_end)
                top_cell = ws.cell(row=h1_row, column=c_start)
                top_cell.value = text
                top_cell.font = font
                top_cell.alignment = align_center
                for r in (h1_row, h2_row):
                    c = ws.cell(row=r, column=c_start)
                    c.fill = fill
                    c.border = thin_border
            else:
                # Horizontal merge across 3 cols
                ws.merge_cells(start_row=h1_row, start_column=c_start, end_row=h1_row, end_column=c_end)
                top_cell = ws.cell(row=h1_row, column=c_start)
                top_cell.value = text
                top_cell.font = font
                top_cell.alignment = align_center
                for col in range(c_start, c_end + 1):
                    c = ws.cell(row=h1_row, column=col)
                    c.fill = fill
                    c.border = thin_border

        # Header 2 sub-headers
        sub_headers = [
            # 1..4 merged with top
            (5, "Tồn", fill_olive, font_header_black),
            (6, "Quá hạn", fill_olive, font_header_black),
            (7, "Quá hạn 30N", fill_olive, font_header_black),
            (8, "Tồn", fill_peach, font_header_black),
            (9, "Quá hạn", fill_peach, font_header_black),
            (10, "Quá hạn 30N", fill_peach, font_header_black),
            (11, "Tồn", fill_aqua, font_header_black),
            (12, "Quá hạn", fill_aqua, font_header_black),
            (13, "Quá hạn 30N", fill_aqua, font_header_black),
            (14, "Tồn", fill_dark_blue, font_header_white),
            (15, "Quá hạn", fill_dark_blue, font_header_white),
            (16, "Quá hạn 30N", fill_dark_blue, font_header_white),
            (17, "Tồn", fill_green, font_header_black),
            (18, "Quá hạn", fill_green, font_header_black),
            (19, "Quá hạn 30N", fill_green, font_header_black),
            # 20 merged with top
        ]

        for col_idx, text, fill, font in sub_headers:
            c = ws.cell(row=h2_row, column=col_idx)
            c.value = text
            c.font = font
            c.fill = fill
            c.alignment = align_center
            c.border = thin_border

        # 4. DATA ROWS
        current_row = h2_row + 1
        for emp in center.employees:
            ws.row_dimensions[current_row].height = 19
            is_tagged = emp.is_tagged

            row_values = [
                emp.username,
                emp.full_name,
                emp.center_raw,
                emp.position,
                emp.icms_bd_nt_ton if emp.icms_bd_nt_ton > 0 else "",
                emp.icms_bd_nt_qh if emp.icms_bd_nt_qh > 0 else "",
                emp.icms_bd_nt_qh30 if emp.icms_bd_nt_qh30 > 0 else "",
                emp.icms_bd_thc_ton if emp.icms_bd_thc_ton > 0 else "",
                emp.icms_bd_thc_qh if emp.icms_bd_thc_qh > 0 else "",
                emp.icms_bd_thc_qh30 if emp.icms_bd_thc_qh30 > 0 else "",
                emp.icms_tt_tc_ton if emp.icms_tt_tc_ton > 0 else "",
                emp.icms_tt_tc_qh if emp.icms_tt_tc_qh > 0 else "",
                emp.icms_tt_tc_qh30 if emp.icms_tt_tc_qh30 > 0 else "",
                emp.bao_hong_ton if emp.bao_hong_ton > 0 else "",
                emp.bao_hong_qh if emp.bao_hong_qh > 0 else "",
                emp.bao_hong_qh30 if emp.bao_hong_qh30 > 0 else "",
                emp.ton if emp.ton > 0 else "",
                emp.qua_han if emp.qua_han > 0 else "",
                emp.qua_han_30 if emp.qua_han_30 > 0 else "",
                emp.thuc_hien_ngay if emp.thuc_hien_ngay > 0 else ""
            ]

            for c_idx, val in enumerate(row_values, start=1):
                cell = ws.cell(row=current_row, column=c_idx, value=val)
                cell.border = thin_border

                # Alignment & Font
                if c_idx == 1:
                    cell.alignment = align_left
                    cell.font = font_data_bold
                elif c_idx in (2, 3):
                    cell.alignment = align_left
                    cell.font = font_data
                elif c_idx == 4:
                    cell.alignment = align_center
                    cell.font = font_data
                elif c_idx == 17: # Cột Q: Tồn
                    cell.alignment = align_center
                    cell.font = font_data_tagged if is_tagged else font_data_bold
                elif c_idx == 20: # Cột T: Thực hiện ngày
                    cell.alignment = align_center
                    cell.font = font_data_bold
                else:
                    cell.alignment = align_center
                    cell.font = font_data

            current_row += 1

        # 5. TOTAL / SUMMARY ROW AT BOTTOM
        ws.row_dimensions[current_row].height = 22
        subtotals = center.get_subtotal()

        # Merge A..D for "Tổng cộng"
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
        c_tot = ws.cell(row=current_row, column=1)
        c_tot.value = f"Tổng cộng ({len(center.employees)} FT - {center.tagged_count} Tag):"
        c_tot.font = font_data_bold
        c_tot.alignment = Alignment(horizontal="right", vertical="center")

        total_vals = [
            (5, subtotals["icms_bd_nt_ton"]),
            (6, subtotals["icms_bd_nt_qh"]),
            (7, subtotals["icms_bd_nt_qh30"]),
            (8, subtotals["icms_bd_thc_ton"]),
            (9, subtotals["icms_bd_thc_qh"]),
            (10, subtotals["icms_bd_thc_qh30"]),
            (11, subtotals["icms_tt_tc_ton"]),
            (12, subtotals["icms_tt_tc_qh"]),
            (13, subtotals["icms_tt_tc_qh30"]),
            (14, subtotals["bao_hong_ton"]),
            (15, subtotals["bao_hong_qh"]),
            (16, subtotals["bao_hong_qh30"]),
            (17, subtotals["ton"]),
            (18, subtotals["qua_han"]),
            (19, subtotals["qua_han_30"]),
            (20, subtotals["thuc_hien_ngay"]),
        ]

        for col_idx in range(1, 21):
            c = ws.cell(row=current_row, column=col_idx)
            c.fill = fill_total
            c.border = double_bottom_border

        for col_idx, val in total_vals:
            c = ws.cell(row=current_row, column=col_idx, value=val if val > 0 else "")
            c.font = font_data_bold
            c.alignment = align_center

        # =========================================================================
        # 6. WO TABLE SECTION (BẢNG WO QUÁ HẠN PHÍA DƯỚI ĐỂ COPY)
        # =========================================================================
        # 6.1 Spacing row
        current_row += 1
        ws.row_dimensions[current_row].height = 14

        # 6.2 Section Title Banner
        current_row += 1
        ws.row_dimensions[current_row].height = 26
        wo_count = len(center.work_orders)
        wo_title_text = f"DANH SÁCH CHI TIẾT WO QUÁ HẠN - {center.center_name.upper()} ({wo_count} WO)"

        # Merge across columns A to O (15 columns matching the WO table)
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=15)
        c_wo_title = ws.cell(row=current_row, column=1)
        c_wo_title.value = wo_title_text
        c_wo_title.font = Font(name=self.font_family, size=11, bold=True, color="FFFFFF")
        c_wo_title.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        fill_wo_title = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        for col_idx in range(1, 16):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = fill_wo_title
            cell.border = thin_border

        if wo_count == 0:
            current_row += 1
            ws.row_dimensions[current_row].height = 22
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=15)
            c_empty = ws.cell(row=current_row, column=1)
            c_empty.value = f"✓ Trung tâm {center.center_name} hiện không có WO nào quá hạn."
            c_empty.font = Font(name=self.font_family, size=10, italic=True, color="2E7D32")
            c_empty.alignment = Alignment(horizontal="center", vertical="center")
            for col_idx in range(1, 16):
                ws.cell(row=current_row, column=col_idx).border = thin_border
        else:
            # 6.3 WO Table Headers
            current_row += 1
            ws.row_dimensions[current_row].height = 24
            wo_headers = [
                (1, "STT"),
                (2, "Mã công việc"),
                (3, "Loại công việc"),
                (4, "Nội dung công việc"),
                (5, "Thời điểm kết thúc"),
                (6, "Trạng thái"),
                (7, "Nhân viên thực hiện"),
                (8, "Mã trạm"),
                (9, "FT comment"),
                (10, "FT mobile"),
                (11, "Ngày hoàn thành"),
                (12, "Ngày quá hạn"),
                (13, "TTH"),
                (14, "BKK"),
                (15, "Lý do BKK"),
            ]

            fill_wo_header = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid") # Dark navy blue
            font_wo_header = Font(name=self.font_family, size=9.5, bold=True, color="FFFFFF")

            for col_idx, h_name in wo_headers:
                c = ws.cell(row=current_row, column=col_idx, value=h_name)
                c.fill = fill_wo_header
                c.font = font_wo_header
                c.alignment = align_center
                c.border = thin_border

            # 6.4 WO Data Rows
            font_wo_data = Font(name=self.font_family, size=9.5, color="000000")
            font_wo_bold = Font(name=self.font_family, size=9.5, bold=True, color="000000")
            font_wo_overdue = Font(name=self.font_family, size=9.5, bold=True, color="C00000")
            fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

            for idx, wo in enumerate(center.work_orders):
                current_row += 1
                ws.row_dimensions[current_row].height = 20

                # Format overdue days
                overdue_val = wo.overdue_days
                if isinstance(overdue_val, float):
                    overdue_display = round(overdue_val, 1)
                else:
                    overdue_display = overdue_val if overdue_val is not None else ""

                row_vals = [
                    idx + 1,
                    wo.code,
                    wo.job_type,
                    wo.content,
                    wo.deadline,
                    wo.status,
                    wo.performer,
                    wo.station_code,
                    wo.ft_comment or "",
                    wo.ft_mobile or "",
                    wo.completed_date or "",
                    overdue_display,
                    wo.center_raw,
                    wo.bkk or "",
                    wo.bkk_reason or ""
                ]

                # Zebra striping for readability
                use_zebra = (idx % 2 == 1)

                for col_idx, val in enumerate(row_vals, start=1):
                    cell = ws.cell(row=current_row, column=col_idx, value=val)
                    cell.border = thin_border
                    if use_zebra:
                        cell.fill = fill_zebra

                    # Alignment & Font
                    if col_idx == 1:
                        cell.alignment = align_center
                        cell.font = font_wo_data
                    elif col_idx == 2:
                        cell.alignment = align_left
                        cell.font = font_wo_bold
                    elif col_idx in (3, 4, 9, 13, 15):
                        cell.alignment = align_left
                        cell.font = font_wo_data
                    elif col_idx == 7: # Performer
                        cell.alignment = align_left
                        cell.font = font_wo_bold
                    elif col_idx == 12: # Overdue
                        cell.alignment = align_center
                        cell.font = font_wo_overdue
                        if isinstance(val, (int, float)):
                            cell.number_format = "0.0"
                    else:
                        cell.alignment = align_center
                        cell.font = font_wo_data

            # 6.5 WO Summary / Footer row
            current_row += 1
            ws.row_dimensions[current_row].height = 22
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=11)
            c_wo_tot = ws.cell(row=current_row, column=1)
            c_wo_tot.value = f"Tổng cộng WO quá hạn ({center.center_name}):"
            c_wo_tot.font = font_data_bold
            c_wo_tot.alignment = Alignment(horizontal="right", vertical="center")

            c_wo_cnt = ws.cell(row=current_row, column=12, value=wo_count)
            c_wo_cnt.font = font_data_tagged
            c_wo_cnt.alignment = align_center

            fill_wo_summary = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Soft yellow
            for col_idx in range(1, 16):
                c = ws.cell(row=current_row, column=col_idx)
                c.fill = fill_wo_summary
                c.border = double_bottom_border

        # =========================================================================
        # 7. PENDING WO TABLE SECTION (BẢNG WO TỒN CHƯA QUÁ HẠN PHÍA DƯỚI - CỘT L > 0)
        # =========================================================================
        # 7.1 Spacing row
        current_row += 1
        ws.row_dimensions[current_row].height = 14

        # 7.2 Section Title Banner
        current_row += 1
        ws.row_dimensions[current_row].height = 26
        pending_count = len(center.pending_work_orders)
        pending_title_text = f"DANH SÁCH CHI TIẾT WO TỒN CHƯA QUÁ HẠN - {center.center_name.upper()} ({pending_count} WO)"

        # Merge across columns A to O (15 columns matching the WO table)
        ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=15)
        c_pending_title = ws.cell(row=current_row, column=1)
        c_pending_title.value = pending_title_text
        c_pending_title.font = Font(name=self.font_family, size=11, bold=True, color="FFFFFF")
        c_pending_title.alignment = Alignment(horizontal="left", vertical="center", indent=1)

        fill_pending_title = PatternFill(start_color="135E48", end_color="135E48", fill_type="solid") # Deep Forest Teal
        for col_idx in range(1, 16):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = fill_pending_title
            cell.border = thin_border

        if pending_count == 0:
            current_row += 1
            ws.row_dimensions[current_row].height = 22
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=15)
            c_empty_pending = ws.cell(row=current_row, column=1)
            c_empty_pending.value = f"✓ Trung tâm {center.center_name} hiện không có WO tồn chưa quá hạn."
            c_empty_pending.font = Font(name=self.font_family, size=10, italic=True, color="135E48")
            c_empty_pending.alignment = Alignment(horizontal="center", vertical="center")
            for col_idx in range(1, 16):
                ws.cell(row=current_row, column=col_idx).border = thin_border
        else:
            # 7.3 Pending WO Table Headers
            current_row += 1
            ws.row_dimensions[current_row].height = 24
            pending_headers = [
                (1, "STT"),
                (2, "Mã công việc"),
                (3, "Loại công việc"),
                (4, "Nội dung công việc"),
                (5, "Thời điểm kết thúc"),
                (6, "Trạng thái"),
                (7, "Nhân viên thực hiện"),
                (8, "Mã trạm"),
                (9, "FT comment"),
                (10, "FT mobile"),
                (11, "Ngày hoàn thành"),
                (12, "Ngày quá hạn"),
                (13, "TTH"),
                (14, "BKK"),
                (15, "Lý do BKK"),
            ]

            fill_pending_header = PatternFill(start_color="1E8449", end_color="1E8449", fill_type="solid") # Elegant Green
            font_pending_header = Font(name=self.font_family, size=9.5, bold=True, color="FFFFFF")

            for col_idx, h_name in pending_headers:
                c = ws.cell(row=current_row, column=col_idx, value=h_name)
                c.fill = fill_pending_header
                c.font = font_pending_header
                c.alignment = align_center
                c.border = thin_border

            # 7.4 Pending WO Data Rows
            font_pending_data = Font(name=self.font_family, size=9.5, color="000000")
            font_pending_bold = Font(name=self.font_family, size=9.5, bold=True, color="000000")
            font_pending_days = Font(name=self.font_family, size=9.5, bold=True, color="135E48") # Dark Green
            fill_pending_zebra = PatternFill(start_color="F4FBF7", end_color="F4FBF7", fill_type="solid")

            for idx, wo in enumerate(center.pending_work_orders):
                current_row += 1
                ws.row_dimensions[current_row].height = 20

                # Format overdue days (>0)
                overdue_val = wo.overdue_days
                if isinstance(overdue_val, float):
                    overdue_display = round(overdue_val, 1)
                else:
                    overdue_display = overdue_val if overdue_val is not None else ""

                row_vals = [
                    idx + 1,
                    wo.code,
                    wo.job_type,
                    wo.content,
                    wo.deadline,
                    wo.status,
                    wo.performer,
                    wo.station_code,
                    wo.ft_comment or "",
                    wo.ft_mobile or "",
                    wo.completed_date or "",
                    overdue_display,
                    wo.center_raw,
                    wo.bkk or "",
                    wo.bkk_reason or ""
                ]

                use_zebra = (idx % 2 == 1)

                for col_idx, val in enumerate(row_vals, start=1):
                    cell = ws.cell(row=current_row, column=col_idx, value=val)
                    cell.border = thin_border
                    if use_zebra:
                        cell.fill = fill_pending_zebra

                    if col_idx == 1:
                        cell.alignment = align_center
                        cell.font = font_pending_data
                    elif col_idx == 2:
                        cell.alignment = align_left
                        cell.font = font_pending_bold
                    elif col_idx in (3, 4, 9, 13, 15):
                        cell.alignment = align_left
                        cell.font = font_pending_data
                    elif col_idx == 7: # Performer
                        cell.alignment = align_left
                        cell.font = font_pending_bold
                    elif col_idx == 12: # Overdue days (>0)
                        cell.alignment = align_center
                        cell.font = font_pending_days
                        if isinstance(val, (int, float)):
                            cell.number_format = "0.0"
                    else:
                        cell.alignment = align_center
                        cell.font = font_pending_data

            # 7.5 Pending WO Summary / Footer row
            current_row += 1
            ws.row_dimensions[current_row].height = 22
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=11)
            c_p_tot = ws.cell(row=current_row, column=1)
            c_p_tot.value = f"Tổng cộng WO tồn chưa quá hạn ({center.center_name}):"
            c_p_tot.font = font_data_bold
            c_p_tot.alignment = Alignment(horizontal="right", vertical="center")

            c_p_cnt = ws.cell(row=current_row, column=12, value=pending_count)
            c_p_cnt.font = Font(name=self.font_family, size=10, bold=True, color="135E48")
            c_p_cnt.alignment = align_center

            fill_pending_summary = PatternFill(start_color="D5F5E3", end_color="D5F5E3", fill_type="solid") # Soft green pastel
            for col_idx in range(1, 16):
                c = ws.cell(row=current_row, column=col_idx)
                c.fill = fill_pending_summary
                c.border = double_bottom_border

        # 8. COLUMN WIDTHS OPTIMIZATION (Hài hòa cho cả bảng FT và bảng WO)
        col_widths = {
            "A": 13,   # FT Username / STT
            "B": 28,   # Tên FT / Mã công việc (WO_ICMS_...)
            "C": 28,   # TTH / Loại công việc
            "D": 26,   # Vị trí / Nội dung công việc
            "E": 19,   # BDNT Tồn / Thời điểm yêu cầu kết thúc
            "F": 12,   # BDNT Quá hạn / Trạng thái
            "G": 16,   # BDNT QH30 / Nhân viên thực hiện
            "H": 11,   # BD THC Tồn / Mã trạm
            "I": 12,   # BD THC Quá hạn / FT comment
            "J": 12,   # BD THC QH30 / FT mobile
            "K": 13,   # TT TC Tồn / Ngày hoàn thành
            "L": 12,   # TT TC Quá hạn / Ngày quá hạn
            "M": 26,   # TT TC QH30 / TTH
            "N": 10,   # Báo hỏng Tồn / BKK
            "O": 16,   # Báo hỏng Quá hạn / Lý do BKK
            "P": 8.5,  # Báo hỏng QH30
            "Q": 8.5,  # Tổng Tồn
            "R": 8.5,  # Tổng Quá hạn
            "S": 9.0,  # Tổng QH30
            "T": 11.5  # Thực hiện ngày
        }

        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width

        # 8. PAGE SETUP FOR PRINTING
        ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True

        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0 # Tự động xuống trang khi có nhiều dòng

        # Narrow margins
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.4
        ws.page_margins.bottom = 0.4
        ws.page_margins.header = 0.2
        ws.page_margins.footer = 0.2

        # Repeat table headers if multiple pages
        ws.print_title_rows = "3:4"
