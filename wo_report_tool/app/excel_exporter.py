"""
Excel exporter module for WO Report Tool.
Exports styled report workbooks containing 'BaoCao' and 'ChiTiet' sheets using openpyxl.
"""

import logging
import os
from datetime import date
from typing import Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.models import ReportData

logger = logging.getLogger(__name__)


class ExcelExporter:
    """Exports structured and beautifully formatted Excel reports."""

    @staticmethod
    def _apply_border(ws, start_row: int, start_col: int, end_row: int, end_col: int, border: Border):
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                ws.cell(row=r, column=c).border = border

    @classmethod
    def export_report(cls, data: ReportData, output_path: str, preview_text: Optional[str] = None) -> str:
        """
        Generates and saves the report Excel file.
        Contains:
          1. Sheet 'BaoCao': Overview report matching Word layout
          2. Sheet 'ChiTiet': Detailed table of all district centers
        """
        logger.info(f"Generating export Excel file at: {output_path}")

        wb = openpyxl.Workbook()
        # Default sheet -> rename to BaoCao
        ws_report = wb.active
        ws_report.title = "BaoCao"

        # Styles
        font_title = Font(name="Arial", size=16, bold=True, color="1E3A8A")
        font_section = Font(name="Arial", size=12, bold=True, color="0F172A")
        font_bold = Font(name="Arial", size=11, bold=True, color="000000")
        font_regular = Font(name="Arial", size=11, color="1E293B")
        font_italic = Font(name="Arial", size=10, italic=True, color="64748B")
        font_highlight_metric = Font(name="Arial", size=13, bold=True, color="DC2626")

        fill_header = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
        fill_sub_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        fill_metric_box = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
        fill_total_row = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")

        align_center = Alignment(horizontal="center", vertical="center")
        align_left = Alignment(horizontal="left", vertical="center")
        align_right = Alignment(horizontal="right", vertical="center")
        align_top_left = Alignment(horizontal="left", vertical="top", wrap_text=True)

        thin_side = Side(border_style="thin", color="CBD5E1")
        thin_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        double_bottom = Border(bottom=Side(border_style="double", color="000000"))

        # ==========================================
        # 1. SHEET 'BaoCao'
        # ==========================================
        ws_report.views.sheetView[0].showGridLines = True

        # Title (A1:G1)
        ws_report.merge_cells("A1:G1")
        cell_a1 = ws_report["A1"]
        cell_a1.value = "BÁO CÁO TIẾN ĐỘ CẮT TỒN WO THÁNG " + str(data.month)
        cell_a1.font = font_title
        cell_a1.alignment = align_center
        ws_report.row_dimensions[1].height = 36

        # Subtitle Date (A2:G2)
        ws_report.merge_cells("A2:G2")
        cell_a2 = ws_report["A2"]
        cell_a2.value = f"Thời điểm báo cáo: Ngày {data.day:02d} tháng {data.month:02d} năm {data.year}"
        cell_a2.font = font_italic
        cell_a2.alignment = align_center
        ws_report.row_dimensions[2].height = 20

        # Recipient (A4:G4)
        ws_report.cell(row=4, column=1, value="Người nhận:").font = font_bold
        ws_report.merge_cells("B4:G4")
        cell_b4 = ws_report["B4"]
        cell_b4.value = data.recipient
        cell_b4.font = font_regular
        cell_b4.alignment = align_left
        ws_report.row_dimensions[4].height = 24

        # Key Metrics Summary Grid (Rows 6-8)
        ws_report.merge_cells("A6:G6")
        ws_report["A6"] = "I. CHỈ SỐ TỔNG HỢP TOÀN TỈNH"
        ws_report["A6"].font = font_section

        # Metric Table Header (Row 7)
        metric_headers = ["Thực hiện ngày", "Tổng tồn WO", "Tồn quá hạn", "Tồn >30 ngày", "Top 5 tồn cao"]
        metric_cols = [("A7:B7", 1), ("C7", 3), ("D7", 4), ("E7", 5), ("F7:G7", 6)]

        ws_report.merge_cells("A7:B7")
        ws_report["A7"] = "Thực hiện ngày"
        ws_report["C7"] = "Tổng tồn WO"
        ws_report["D7"] = "Tồn quá hạn"
        ws_report["E7"] = "Tồn > 30 ngày"
        ws_report.merge_cells("F7:G7")
        ws_report["F7"] = "Top 5 trung tâm tồn cao"

        for col_name in ["A7", "B7", "C7", "D7", "E7", "F7", "G7"]:
            cell = ws_report[col_name]
            cell.font = font_bold
            cell.fill = fill_sub_header
            cell.alignment = align_center
            cell.border = thin_border

        ws_report.row_dimensions[7].height = 25

        # Metric Values (Row 8)
        ws_report.merge_cells("A8:B8")
        ws_report["A8"] = f"{data.thuc_hien_ngay} WO"
        ws_report["C8"] = f"{data.tong_ton} WO"
        ws_report["D8"] = f"{data.ton_qua_han} WO"
        
        detail_30 = ""
        if data.overdue_30_centers:
            detail_30 = " (" + ", ".join(f"{c}:{cnt}" for c, cnt in data.overdue_30_centers) + ")"
        ws_report["E8"] = f"{data.ton_tren_30_ngay} WO{detail_30}"
        
        ws_report.merge_cells("F8:G8")
        ws_report["F8"] = ", ".join(data.top_centers)

        for col_name in ["A8", "B8", "C8", "D8", "E8", "F8", "G8"]:
            cell = ws_report[col_name]
            cell.font = font_highlight_metric if col_name in ["A8", "C8", "D8", "E8"] else font_bold
            cell.fill = fill_metric_box
            cell.alignment = align_center
            cell.border = thin_border

        ws_report.row_dimensions[8].height = 30

        # Section II: Message Body (Row 10+)
        ws_report.merge_cells("A10:G10")
        ws_report["A10"] = "II. NỘI DUNG THÔNG BÁO / ĐIỀU HÀNH (DÙNG ĐỂ GỬI LÊN NHÓM TAMMI)"
        ws_report["A10"].font = font_section

        # Complete formatted message box
        content_to_put = preview_text or ""
        ws_report.merge_cells("A11:G18")
        msg_cell = ws_report["A11"]
        msg_cell.value = content_to_put
        msg_cell.font = font_regular
        msg_cell.alignment = align_top_left
        cls._apply_border(ws_report, 11, 1, 18, 7, thin_border)

        # Column widths for BaoCao sheet
        col_widths_bc = {"A": 16, "B": 18, "C": 16, "D": 16, "E": 20, "F": 16, "G": 22}
        for col_letter, width in col_widths_bc.items():
            ws_report.column_dimensions[col_letter].width = width

        # ==========================================
        # 2. SHEET 'ChiTiet'
        # ==========================================
        ws_detail = wb.create_sheet(title="ChiTiet")
        ws_detail.views.sheetView[0].showGridLines = True

        detail_headers = [
            "STT", "Mã TT", "Tên trung tâm", "Tổng tồn",
            "Tồn quá hạn", "Tồn >30 ngày", "TH Ngày", "KH Ngày",
            "Tỷ lệ HT Ngày", "Đánh giá", "Họp Ngày"
        ]

        ws_detail.row_dimensions[1].height = 28
        for col_num, h_text in enumerate(detail_headers, 1):
            cell = ws_detail.cell(row=1, column=col_num, value=h_text)
            cell.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            cell.fill = fill_header
            cell.alignment = align_center
            cell.border = thin_border

        current_row = 2
        for center in data.centers:
            ws_detail.row_dimensions[current_row].height = 20

            ws_detail.cell(row=current_row, column=1, value=center.stt).alignment = align_center
            ws_detail.cell(row=current_row, column=2, value=center.code).alignment = align_center
            ws_detail.cell(row=current_row, column=3, value=center.full_name).alignment = align_left
            ws_detail.cell(row=current_row, column=4, value=center.tong_ton).alignment = align_right
            ws_detail.cell(row=current_row, column=5, value=center.ton_qua_han).alignment = align_right
            ws_detail.cell(row=current_row, column=6, value=center.ton_30_ngay).alignment = align_right
            ws_detail.cell(row=current_row, column=7, value=center.thuc_hien_ngay).alignment = align_right
            ws_detail.cell(row=current_row, column=8, value=center.kh_ngay).alignment = align_right

            pct_cell = ws_detail.cell(row=current_row, column=9, value=center.ty_le_ht_ngay)
            pct_cell.number_format = "0.0%"
            pct_cell.alignment = align_right

            ws_detail.cell(row=current_row, column=10, value=center.danh_gia).alignment = align_center
            ws_detail.cell(row=current_row, column=11, value=center.hop_ngay).alignment = align_center

            for c in range(1, 12):
                ws_detail.cell(row=current_row, column=c).border = thin_border
                ws_detail.cell(row=current_row, column=c).font = font_regular

            current_row += 1

        # Total Row (Hàng TỔNG CỘNG)
        ws_detail.row_dimensions[current_row].height = 24
        total_cell = ws_detail.cell(row=current_row, column=3, value="TỔNG CỘNG TOÀN TỈNH")
        total_cell.font = font_bold
        total_cell.alignment = align_left

        ws_detail.cell(row=current_row, column=4, value=data.tong_ton).alignment = align_right
        ws_detail.cell(row=current_row, column=5, value=data.ton_qua_han).alignment = align_right
        ws_detail.cell(row=current_row, column=6, value=data.ton_tren_30_ngay).alignment = align_right
        ws_detail.cell(row=current_row, column=7, value=data.thuc_hien_ngay).alignment = align_right

        sum_kh = sum(c.kh_ngay for c in data.centers)
        ws_detail.cell(row=current_row, column=8, value=sum_kh).alignment = align_right

        overall_pct = (data.thuc_hien_ngay / sum_kh) if sum_kh > 0 else 0.0
        tot_pct_cell = ws_detail.cell(row=current_row, column=9, value=overall_pct)
        tot_pct_cell.number_format = "0.0%"
        tot_pct_cell.alignment = align_right

        for c in range(1, 12):
            cell = ws_detail.cell(row=current_row, column=c)
            cell.font = font_bold
            cell.fill = fill_total_row
            cell.border = thin_border

        # Freeze Panes on ChiTiet (Header fixed)
        ws_detail.freeze_panes = "A2"

        # Auto Filter on ChiTiet
        last_col_letter = get_column_letter(len(detail_headers))
        ws_detail.auto_filter.ref = f"A1:{last_col_letter}{current_row - 1}"

        # Auto-adjust column widths for ChiTiet
        for col in ws_detail.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws_detail.column_dimensions[col_letter].width = max(max_len + 4, 12)
        # Give full name column plenty of room
        ws_detail.column_dimensions["C"].width = 44

        # Save workbook
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        wb.save(output_path)
        logger.info(f"Excel report successfully saved to {output_path}")
        return output_path
