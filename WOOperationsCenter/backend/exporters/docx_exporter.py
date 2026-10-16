"""
Word (DOCX) Exporter for WO Operations Center.
Generates styled Microsoft Word documents matching official unit templates.
"""

import logging
import os
from typing import List, Optional

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from backend.domain.models import CenterReport, ProvinceSummary

logger = logging.getLogger(__name__)


class DocxExporter:
    """Exports Word report documents."""

    @classmethod
    def export_center_report_docx(
        cls,
        center: CenterReport,
        notification: str,
        output_path: str
    ) -> str:
        """Exports a single center report to Word DOCX."""
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        doc = docx.Document()

        # Set normal font
        style = doc.styles['Normal']
        font = style.font
        font.name = 'Arial'
        font.size = Pt(11)

        # Title
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run(f"BÁO CÁO ĐIỀU HÀNH WO — TRUNG TÂM {center.center_name.upper()} ({center.center_code})")
        run_title.bold = True
        run_title.font.size = Pt(14)
        run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

        # Notification & TAG
        p_notif = doc.add_paragraph()
        run_notif = p_notif.add_run(notification)
        run_notif.bold = True

        doc.add_paragraph()  # spacing

        # Employee Table
        table = doc.add_table(rows=2, cols=8)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'

        # Headers
        hdr_cells = table.rows[0].cells
        hdr_titles = ["FT", "Tên FT", "TTH", "Vị trí", "Tổng Tồn", "Quá hạn", "QH 30N", "TH Ngày"]
        for i, title in enumerate(hdr_titles):
            hdr_cells[i].text = title
            hdr_cells[i].paragraphs[0].runs[0].bold = True

        for emp in center.employees:
            row_cells = table.add_row().cells
            row_cells[0].text = emp.username
            row_cells[1].text = emp.full_name
            row_cells[2].text = emp.center_code
            row_cells[3].text = emp.position
            row_cells[4].text = str(emp.ton) if emp.ton > 0 else ""
            row_cells[5].text = str(emp.qua_han) if emp.qua_han > 0 else ""
            row_cells[6].text = str(emp.qua_han_30) if emp.qua_han_30 > 0 else ""
            row_cells[7].text = str(emp.thuc_hien_ngay) if emp.thuc_hien_ngay > 0 else ""

        # Subtotal Row
        sub = center.get_subtotal()
        sub_cells = table.add_row().cells
        sub_cells[0].text = "Tổng cộng"
        sub_cells[0].paragraphs[0].runs[0].bold = True
        sub_cells[4].text = str(sub['ton'])
        sub_cells[5].text = str(sub['qua_han'])
        sub_cells[6].text = str(sub['qua_han_30'])
        sub_cells[7].text = str(sub['thuc_hien_ngay'])
        for c in sub_cells:
            if c.paragraphs and c.paragraphs[0].runs:
                c.paragraphs[0].runs[0].bold = True

        doc.save(output_path)
        logger.info(f"DOCX report saved to {output_path}")
        return output_path
